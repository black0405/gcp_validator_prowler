"""Write the consolidated validation report to a standard .xlsx workbook."""
from __future__ import annotations

from typing import Iterable, List

from .models import (
    AGREE, INCONCLUSIVE, LIKELY_FALSE_NEGATIVE, LIKELY_FALSE_POSITIVE,
    NO_PROWLER, CheckResult, Method, ResourceResult, Verdict,
)

COLUMNS = [
    "Check ID", "Service", "Severity", "Title", "Prowler Hub Link",
    "Project", "Region", "Resource Name", "Resource ID",
    "Prowler Status", "Prowler Detail",
    "M1 API", "M1 Detail",
    "M2 Prowler-replica", "M2 Detail",
    "M3 Logs", "M3 Detail",
    "M4 Alternate", "M4 Detail",
    "M5 Exposure", "M5 Detail",
    "Consensus", "Agreement", "Check Error",
]

# Fills for the Agreement column (openpyxl colors, ARGB).
_FILLS = {
    LIKELY_FALSE_POSITIVE: "FFF6C244",   # amber — Prowler FAIL we could not confirm
    LIKELY_FALSE_NEGATIVE: "FFE06666",   # red   — Prowler PASS but we found a problem
    AGREE: "FFB6D7A8",                   # green — verdicts match
    INCONCLUSIVE: "FFD9D9D9",            # grey
    NO_PROWLER: "FFFFFFFF",
}


def _m(rr: ResourceResult, method: Method):
    r = rr.methods.get(method)
    if r is None:
        return "", ""
    return r.verdict.value, r.detail


def _rows_for_check(cr: CheckResult) -> Iterable[List[str]]:
    if not cr.resources:
        # Emit one informational row so a check that found nothing is still visible.
        yield [
            cr.check_id, cr.service, cr.severity, cr.title, cr.hub_link,
            "", "", "", "", "", "",
            "", "", "", "", "", "", "", "", "", "",
            Verdict.NA.value, INCONCLUSIVE if not cr.error else "ERROR", cr.error or "no resources evaluated",
        ]
        return
    for rr in cr.resources:
        m1 = _m(rr, Method.API)
        m2 = _m(rr, Method.PROWLER_REPLICA)
        m3 = _m(rr, Method.LOGS)
        m4 = _m(rr, Method.ALTERNATE)
        m5 = _m(rr, Method.EXPOSURE)
        yield [
            cr.check_id, cr.service, cr.severity, cr.title, cr.hub_link,
            rr.project, rr.region, rr.resource_name, rr.resource_id,
            rr.prowler_status.value if rr.prowler_status else "", rr.prowler_detail,
            m1[0], m1[1], m2[0], m2[1], m3[0], m3[1], m4[0], m4[1], m5[0], m5[1],
            rr.consensus().value, rr.agreement(), cr.error or "",
        ]


def write_report(check_results: List[CheckResult], out_path: str) -> dict:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Validation"

    header_font = Font(bold=True, color="FFFFFFFF")
    header_fill = PatternFill("solid", fgColor="FF434343")
    hub_col = COLUMNS.index("Prowler Hub Link") + 1
    agree_col = COLUMNS.index("Agreement") + 1

    ws.append(COLUMNS)
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    counts = {AGREE: 0, LIKELY_FALSE_POSITIVE: 0, LIKELY_FALSE_NEGATIVE: 0,
              INCONCLUSIVE: 0, NO_PROWLER: 0}

    r = 1
    for cr in sorted(check_results, key=lambda x: (x.service, x.check_id)):
        for row in _rows_for_check(cr):
            r += 1
            ws.append(row)
            agreement = row[agree_col - 1]
            counts[agreement] = counts.get(agreement, 0) + 1
            fill = _FILLS.get(agreement)
            if fill and fill != "FFFFFFFF":
                ws.cell(row=r, column=agree_col).fill = PatternFill("solid", fgColor=fill)
            # Hyperlink the hub column.
            link = row[hub_col - 1]
            if link:
                hc = ws.cell(row=r, column=hub_col)
                hc.hyperlink = link
                hc.font = Font(color="FF1155CC", underline="single")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{r}"

    widths = {
        "Check ID": 42, "Service": 12, "Severity": 9, "Title": 40,
        "Prowler Hub Link": 46, "Project": 18, "Region": 14,
        "Resource Name": 26, "Resource ID": 34, "Prowler Status": 13,
        "Prowler Detail": 40, "Consensus": 11, "Agreement": 22, "Check Error": 30,
    }
    for i, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(name, 30)

    # ---- summary sheet ----
    summ = wb.create_sheet("Summary")
    summ.append(["Metric", "Count"])
    summ.cell(row=1, column=1).font = Font(bold=True)
    summ.cell(row=1, column=2).font = Font(bold=True)
    summ.append(["Checks evaluated", len(check_results)])
    summ.append(["Resource rows", r - 1])
    for k in (AGREE, LIKELY_FALSE_POSITIVE, LIKELY_FALSE_NEGATIVE, INCONCLUSIVE, NO_PROWLER):
        summ.append([k, counts.get(k, 0)])
    summ.column_dimensions["A"].width = 28
    summ.column_dimensions["B"].width = 12

    wb.save(out_path)
    return counts
